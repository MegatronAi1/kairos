"""Impact model: simulated plant history x Kairos's measured improvement rates.

Plant volumes (batches, failures, energy per batch, feedstock) are SIMULATED,
deterministically seeded per plant, because no real plant ledger is connected.
The improvement rates are MEASURED: fault detection rate from
artifacts/metrics/faults.json and the energy saving from artifacts/metrics/energy.json
(held-out batches). Every other number is a labelled assumption.

Formulas (per plant-month, from the go-live month on):
  adoption          = 30%, 60%, 90%, then 90% of batches run with Kairos advice
  rescued batches   ~ Binomial(failed_baseline, detection_rate x rescue_share x adoption)
  energy saved kWh  = batches x kWh/batch x energy_saving x adoption
  feedstock saved   = rescued x feed/batch (not wasted) + batches x feed/batch x feed_opt x adoption
  money             = rescued x product value x recovery fraction
                      + energy saved x tariff + feed-optimisation kg x feed price
(the feedstock of a rescued batch is counted in kg but not again in money: it is
part of the recovered product value).
"""

from __future__ import annotations

import io
import json
import zlib
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Literal

import numpy as np

from app.replay.scenarios import Plant, Tenant
from kairos_ml.constants import load_constants

Period = Literal["month", "quarter", "year"]

MONTHS = 24
LIVE_INDEX = 12  # month 13 = go-live
ADOPTION_RAMP = (0.3, 0.6, 0.9)
SIMULATED_NOTE = "Simulated plant data, illustrative: plant volumes, failures and costs are generated, not reported by a plant."

# Assumptions (value, unit, label). Measured rates are read from artifacts.
RESCUE_SHARE = 0.5
FEED_KG_PER_BATCH = 18_000.0
FEED_OPT_SHARE = 0.02
FEED_PRICE_INR_PER_KG = 45.0
VALUE_USD_PER_BATCH = 45_000.0
INR_PER_USD = 83.0
BATCH_CYCLE_H = 250.0  # ~230 h batch + turnaround

SUM_KEYS = (
    "batches",
    "failed_before_baseline",
    "failed_actual",
    "batches_rescued",
    "energy_kwh_baseline",
    "energy_kwh_actual",
    "energy_saved_kwh",
    "co2_saved_kg",
    "feedstock_saved_kg",
    "money_saved_inr",
    "money_saved_usd",
)
BREAKDOWN_KEYS = ("rescued_batches", "energy", "feedstock")


@dataclass(frozen=True)
class Rates:
    detection_rate: float
    energy_saving_pct: float
    kwh_per_batch: float
    tariff_inr_per_kwh: float
    grid_kg_per_kwh: float
    recovery_fraction: float
    source: str


def load_rates(artifacts_dir: Path) -> Rates:
    c = load_constants()
    detection, saving, kwh, source = 0.9, 25.9, 31_600.0, "defaults (metrics not found)"
    faults_p, energy_p = (
        artifacts_dir / "metrics" / "faults.json",
        artifacts_dir / "metrics" / "energy.json",
    )
    if faults_p.exists() and energy_p.exists():
        k = json.loads(faults_p.read_text(encoding="utf-8"))["summary"]["kairos"]
        e = json.loads(energy_p.read_text(encoding="utf-8"))
        detection = k["detected"] / k["n_fault_batches"]
        saving = float(e["mean_saving_pct"])
        kwh = float(np.mean([b["kwh_now"] for b in e["per_batch"]]))
        source = "artifacts/metrics (held-out evaluation)"
    return Rates(
        detection_rate=detection,
        energy_saving_pct=saving,
        kwh_per_batch=kwh,
        tariff_inr_per_kwh=c["electricity_tariff_india"].value,
        grid_kg_per_kwh=c["grid_emission_factor_india"].value,
        recovery_fraction=c["ponr_recovery_fraction"].value,
        source=source,
    )


def month_labels(today: date) -> list[str]:
    """The 24 calendar months before the current month, oldest first."""
    idx = today.year * 12 + today.month - 1
    out = []
    for i in range(idx - MONTHS, idx):
        out.append(f"{i // 12}-{i % 12 + 1:02d}")
    return out


def adoption(month_index: int) -> float:
    k = month_index - LIVE_INDEX
    if k < 0:
        return 0.0
    return ADOPTION_RAMP[min(k, len(ADOPTION_RAMP) - 1)]


def simulate_plant(plant: Plant, months: list[str], rates: Rates) -> list[dict[str, Any]]:
    rng = np.random.default_rng(zlib.crc32(plant.id.encode()))
    reactors = max(1, len(plant.slots))
    fail_rate = rng.uniform(0.08, 0.12)
    kwh_factor = rng.uniform(0.9, 1.1)
    per_reactor = 720.0 / BATCH_CYCLE_H
    p_rescue = rates.detection_rate * RESCUE_SHARE
    rows = []
    for i, label in enumerate(months):
        batches = max(1, round(reactors * per_reactor + rng.normal(0, 0.8)))
        failed = int(rng.binomial(batches, fail_rate))
        a = adoption(i)
        rescued = int(rng.binomial(failed, p_rescue * a)) if a > 0 else 0
        kwh_base = batches * rates.kwh_per_batch * kwh_factor
        saved_kwh = kwh_base * rates.energy_saving_pct / 100 * a
        feed_opt_kg = batches * FEED_KG_PER_BATCH * FEED_OPT_SHARE * a
        money_rescue = rescued * VALUE_USD_PER_BATCH * INR_PER_USD * rates.recovery_fraction
        money_energy = saved_kwh * rates.tariff_inr_per_kwh
        money_feed = feed_opt_kg * FEED_PRICE_INR_PER_KG
        money = money_rescue + money_energy + money_feed
        rows.append(
            {
                "period": label,
                "batches": batches,
                "failed_before_baseline": failed,
                "failed_actual": failed - rescued,
                "batches_rescued": rescued,
                "energy_kwh_baseline": kwh_base,
                "energy_kwh_actual": kwh_base - saved_kwh,
                "energy_saved_kwh": saved_kwh,
                "co2_saved_kg": saved_kwh * rates.grid_kg_per_kwh,
                "feedstock_saved_kg": rescued * FEED_KG_PER_BATCH + feed_opt_kg,
                "money_saved_inr": money,
                "money_saved_usd": money / INR_PER_USD,
                "breakdown_inr": {
                    "rescued_batches": money_rescue,
                    "energy": money_energy,
                    "feedstock": money_feed,
                },
            }
        )
    return rows


def _empty(period: str) -> dict[str, Any]:
    row: dict[str, Any] = {k: 0 for k in SUM_KEYS}
    row["period"] = period
    row["breakdown_inr"] = {k: 0.0 for k in BREAKDOWN_KEYS}
    return row


def _add(into: dict[str, Any], row: dict[str, Any]) -> None:
    for k in SUM_KEYS:
        into[k] += row[k]
    for k in BREAKDOWN_KEYS:
        into["breakdown_inr"][k] += row["breakdown_inr"][k]


def period_key(month: str, period: Period) -> str:
    y, m = month.split("-")
    if period == "year":
        return y
    if period == "quarter":
        return f"{y}-Q{(int(m) - 1) // 3 + 1}"
    return month


def _round(row: dict[str, Any]) -> dict[str, Any]:
    out = {
        k: (round(v, 1) if isinstance(v, float) else v)
        for k, v in row.items()
        if k != "breakdown_inr"
    }
    out["breakdown_inr"] = {k: round(v, 1) for k, v in row["breakdown_inr"].items()}
    return out


def build_impact(
    tenant: Tenant,
    artifacts_dir: Path,
    period: Period = "month",
    plant: str = "all",
    today: date | None = None,
) -> dict[str, Any]:
    rates = load_rates(artifacts_dir)
    months = month_labels(today or date.today())
    per_plant = {p.id: simulate_plant(p, months, rates) for p in tenant.plants}
    selected = list(per_plant) if plant == "all" else [plant]
    if any(pid not in per_plant for pid in selected):
        raise KeyError(plant)

    series: dict[str, dict[str, Any]] = {}
    totals = _empty("with Kairos")
    baseline_year = _empty("before Kairos")
    for pid in selected:
        for i, row in enumerate(per_plant[pid]):
            key = period_key(row["period"], period)
            _add(series.setdefault(key, _empty(key)), row)
            _add(totals if i >= LIVE_INDEX else baseline_year, row)

    by_plant = []
    for p in tenant.plants:
        live = per_plant[p.id][LIVE_INDEX:]
        by_plant.append(
            {
                "plant_id": p.id,
                "name": p.name,
                "city": p.city,
                "money_saved_inr": round(sum(r["money_saved_inr"] for r in live), 1),
                "money_saved_usd": round(sum(r["money_saved_usd"] for r in live), 1),
                "energy_saved_kwh": round(sum(r["energy_saved_kwh"] for r in live), 1),
                "feedstock_saved_kg": round(sum(r["feedstock_saved_kg"] for r in live), 1),
                "batches_rescued": sum(r["batches_rescued"] for r in live),
            }
        )

    return {
        "tenant": tenant.name,
        "period": period,
        "plant": plant,
        "kairos_live_from": months[LIVE_INDEX],
        "months": [months[0], months[-1]],
        "provenance": {
            "simulated": SIMULATED_NOTE,
            "measured_rates": {
                "detection_rate": round(rates.detection_rate, 3),
                "energy_saving_pct": round(rates.energy_saving_pct, 1),
                "source": rates.source,
            },
            "assumptions": assumptions(rates),
        },
        "series": [_round(series[k]) for k in sorted(series)],
        "totals": _round(totals),
        "baseline_year": _round(baseline_year),
        "by_plant": by_plant,
    }


def assumptions(rates: Rates) -> list[dict[str, Any]]:
    return [
        {
            "key": "detection_rate",
            "label": "Fault batches detected by Kairos",
            "value": round(rates.detection_rate, 3),
            "unit": "-",
            "provenance": "measured",
        },
        {
            "key": "energy_saving_pct",
            "label": "Agitation + aeration energy saving",
            "value": round(rates.energy_saving_pct, 1),
            "unit": "%",
            "provenance": "measured",
        },
        {
            "key": "kwh_per_batch",
            "label": "Agitation + aeration energy per batch (current practice)",
            "value": round(rates.kwh_per_batch),
            "unit": "kWh",
            "provenance": "measured",
        },
        {
            "key": "rescue_share",
            "label": "Detected faults that operators act on in time",
            "value": RESCUE_SHARE,
            "unit": "-",
            "provenance": "assumed",
        },
        {
            "key": "recovery_fraction",
            "label": "Product value kept in a rescued batch",
            "value": rates.recovery_fraction,
            "unit": "-",
            "provenance": "assumed",
        },
        {
            "key": "adoption",
            "label": "Batches run with Kairos advice (months 1, 2, 3+)",
            "value": "30% / 60% / 90%",
            "unit": "-",
            "provenance": "assumed",
        },
        {
            "key": "failure_rate",
            "label": "Batch failure rate before Kairos (per plant)",
            "value": "8-12%",
            "unit": "-",
            "provenance": "simulated",
        },
        {
            "key": "value_per_batch",
            "label": "Product value per batch",
            "value": VALUE_USD_PER_BATCH,
            "unit": "USD",
            "provenance": "assumed",
        },
        {
            "key": "feed_per_batch",
            "label": "Sugar feedstock per batch",
            "value": FEED_KG_PER_BATCH,
            "unit": "kg",
            "provenance": "assumed",
        },
        {
            "key": "feed_opt",
            "label": "Feed saved by tighter feeding",
            "value": FEED_OPT_SHARE * 100,
            "unit": "%",
            "provenance": "assumed",
        },
        {
            "key": "feed_price",
            "label": "Feedstock price",
            "value": FEED_PRICE_INR_PER_KG,
            "unit": "INR/kg",
            "provenance": "assumed",
        },
        {
            "key": "tariff",
            "label": "Electricity tariff",
            "value": rates.tariff_inr_per_kwh,
            "unit": "INR/kWh",
            "provenance": "assumed",
        },
        {
            "key": "grid",
            "label": "Grid emission factor",
            "value": rates.grid_kg_per_kwh,
            "unit": "kg CO2/kWh",
            "provenance": "assumed",
        },
        {
            "key": "fx",
            "label": "Exchange rate",
            "value": INR_PER_USD,
            "unit": "INR/USD",
            "provenance": "assumed",
        },
    ]


# ---------------------------------------------------------------- exports

SERIES_COLUMNS = (
    ("period", "Period", None),
    ("batches", "Batches", "#,##0"),
    ("failed_before_baseline", "Failures without Kairos", "#,##0"),
    ("failed_actual", "Failures", "#,##0"),
    ("batches_rescued", "Batches rescued", "#,##0"),
    ("energy_kwh_baseline", "Energy baseline (kWh)", "#,##0"),
    ("energy_kwh_actual", "Energy actual (kWh)", "#,##0"),
    ("energy_saved_kwh", "Energy saved (kWh)", "#,##0"),
    ("co2_saved_kg", "CO2 avoided (kg)", "#,##0"),
    ("feedstock_saved_kg", "Feedstock saved (kg)", "#,##0"),
    ("money_saved_inr", "Saved (INR)", "#,##0"),
    ("money_saved_usd", "Saved (USD)", "#,##0"),
)
PERIOD_NAMES = {"month": "Monthly", "quarter": "Quarterly", "year": "Yearly"}


def _scope(data: dict[str, Any]) -> str:
    return (
        "All plants"
        if data["plant"] == "all"
        else next(
            (p["name"] for p in data["by_plant"] if p["plant_id"] == data["plant"]), data["plant"]
        )
    )


def to_xlsx(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F766E")

    def header(ws: Any, row: int, titles: list[str]) -> None:
        for col, title in enumerate(titles, start=1):
            cell = ws.cell(row=row, column=col, value=title)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Kairos impact report"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"{data['tenant']} - {_scope(data)} - generated {date.today().isoformat()}"
    ws["A3"] = data["provenance"]["simulated"]
    ws["A3"].font = Font(italic=True, color="B45309")
    ws["A4"] = (
        f"Kairos live from {data['kairos_live_from']}; totals cover the 12 months with Kairos."
    )
    t = data["totals"]
    rows = [
        ("Money saved (INR)", t["money_saved_inr"], "#,##0"),
        ("Money saved (USD)", t["money_saved_usd"], "#,##0"),
        ("  from rescued batches (INR)", t["breakdown_inr"]["rescued_batches"], "#,##0"),
        ("  from energy (INR)", t["breakdown_inr"]["energy"], "#,##0"),
        ("  from feedstock (INR)", t["breakdown_inr"]["feedstock"], "#,##0"),
        ("Energy saved (kWh)", t["energy_saved_kwh"], "#,##0"),
        ("CO2 avoided (kg)", t["co2_saved_kg"], "#,##0"),
        ("Feedstock saved (kg)", t["feedstock_saved_kg"], "#,##0"),
        ("Batches rescued", t["batches_rescued"], "#,##0"),
        ("Batches run", t["batches"], "#,##0"),
    ]
    header(ws, 6, ["Metric (12 months with Kairos)", "Value"])
    for i, (label, value, fmt) in enumerate(rows, start=7):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value).number_format = fmt
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18

    ss = wb.create_sheet(PERIOD_NAMES[data["period"]])
    header(ss, 1, [c[1] for c in SERIES_COLUMNS])
    for r, row in enumerate(data["series"], start=2):
        for c, (key, _, num_fmt) in enumerate(SERIES_COLUMNS, start=1):
            cell = ss.cell(row=r, column=c, value=row[key])
            if num_fmt:
                cell.number_format = num_fmt
    for c in range(1, len(SERIES_COLUMNS) + 1):
        ss.column_dimensions[ss.cell(row=1, column=c).column_letter].width = 16
    ss.row_dimensions[1].height = 32
    ss.freeze_panes = "B2"
    n = len(data["series"])
    if n:
        chart = BarChart()
        chart.title = "Money saved per period (INR)"
        chart.y_axis.title = "INR"
        money_col = [k for k, _, _ in SERIES_COLUMNS].index("money_saved_inr") + 1
        chart.add_data(
            Reference(ss, min_col=money_col, min_row=1, max_row=n + 1), titles_from_data=True
        )
        chart.set_categories(Reference(ss, min_col=1, min_row=2, max_row=n + 1))
        chart.height, chart.width = 9, 22
        ss.add_chart(chart, f"A{n + 4}")

    bp = wb.create_sheet("By plant")
    plant_cols = (
        "Plant",
        "City",
        "Saved (INR)",
        "Saved (USD)",
        "Energy saved (kWh)",
        "Feedstock saved (kg)",
        "Batches rescued",
    )
    header(bp, 1, list(plant_cols))
    for r, p in enumerate(data["by_plant"], start=2):
        vals = (
            p["name"],
            p["city"],
            p["money_saved_inr"],
            p["money_saved_usd"],
            p["energy_saved_kwh"],
            p["feedstock_saved_kg"],
            p["batches_rescued"],
        )
        for c, v in enumerate(vals, start=1):
            cell = bp.cell(row=r, column=c, value=v)
            if c > 2:
                cell.number_format = "#,##0"
    for letter, width in zip("ABCDEFG", (34, 14, 16, 16, 18, 18, 14), strict=True):
        bp.column_dimensions[letter].width = width

    asm = wb.create_sheet("Assumptions")
    header(asm, 1, ["Assumption", "Value", "Unit", "Provenance"])
    for r, a in enumerate(data["provenance"]["assumptions"], start=2):
        for c, v in enumerate((a["label"], a["value"], a["unit"], a["provenance"]), start=1):
            asm.cell(row=r, column=c, value=v)
    asm.cell(
        row=len(data["provenance"]["assumptions"]) + 3,
        column=1,
        value=data["provenance"]["simulated"],
    )
    for letter, width in zip("ABCD", (56, 16, 14, 14), strict=True):
        asm.column_dimensions[letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt(v: float) -> str:
    return f"{v:,.0f}"


def to_pdf(data: dict[str, Any]) -> bytes:
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    teal = colors.HexColor("#0F766E")
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=8, leading=10)
    note = styles["BodyText"].clone(
        "note", fontSize=8.5, leading=11, textColor=colors.HexColor("#B45309")
    )
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="Kairos impact report",
        author="Kairos",
    )
    t = data["totals"]
    story: list[Any] = [
        Paragraph("Kairos impact report", styles["Title"]),
        Paragraph(
            f"{data['tenant']} &middot; {_scope(data)} &middot; {PERIOD_NAMES[data['period']]} &middot; "
            f"generated {date.today().isoformat()} &middot; Kairos live from {data['kairos_live_from']}",
            styles["BodyText"],
        ),
        Paragraph(
            data["provenance"]["simulated"]
            + " Improvement rates are measured on held-out batches.",
            note,
        ),
        Spacer(1, 6),
    ]
    kpis = [
        [
            "Money saved (12 months)",
            "Energy saved",
            "CO2 avoided",
            "Feedstock saved",
            "Batches rescued",
        ],
        [
            f"INR {_fmt(t['money_saved_inr'])}\nUSD {_fmt(t['money_saved_usd'])}",
            f"{t['energy_saved_kwh'] / 1000:,.0f} MWh",
            f"{t['co2_saved_kg'] / 1000:,.0f} t",
            f"{t['feedstock_saved_kg'] / 1000:,.0f} t",
            str(t["batches_rescued"]),
        ],
    ]
    kt = Table(kpis, colWidths=[52 * mm] * 5)
    kt.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), teal),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 1), (-1, 1), 12),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.grey),
                ("TOPPADDING", (0, 1), (-1, 1), 8),
                ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
            ]
        )
    )
    story += [kt, Spacer(1, 8)]

    series = data["series"]
    if series:
        d = Drawing(260 * mm, 62 * mm)
        bc = VerticalBarChart()
        bc.x, bc.y, bc.width, bc.height = 18 * mm, 10 * mm, 236 * mm, 46 * mm
        bc.data = [
            [r["breakdown_inr"]["rescued_batches"] / 1e5 for r in series],
            [r["breakdown_inr"]["energy"] / 1e5 for r in series],
            [r["breakdown_inr"]["feedstock"] / 1e5 for r in series],
        ]
        bc.categoryAxis.style = "stacked"
        bc.categoryAxis.categoryNames = [r["period"] for r in series]
        bc.categoryAxis.labels.fontSize = 6.5
        bc.categoryAxis.labels.angle = 30 if len(series) > 12 else 0
        bc.categoryAxis.labels.dy = -6
        bc.valueAxis.valueMin = 0
        bc.valueAxis.labels.fontSize = 7
        for i, col in enumerate(("#0F766E", "#2563EB", "#D97706")):
            bc.bars[i].fillColor = colors.HexColor(col)
            bc.bars[i].strokeColor = None
        d.add(bc)
        from reportlab.graphics.shapes import String

        d.add(
            String(
                18 * mm,
                58 * mm,
                "Savings per period, INR lakh (teal: rescued batches, blue: energy, amber: feedstock)",
                fontSize=8,
            )
        )
        story += [d, Spacer(1, 6)]

    cols = SERIES_COLUMNS
    table = [[c[1] for c in cols]] + [
        [row[k] if k == "period" else _fmt(row[k]) for k, _, _ in cols] for row in series
    ]
    st = Table(table, repeatRows=1)
    st.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), teal),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 6.8),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
            ]
        )
    )
    story += [
        Paragraph(f"{PERIOD_NAMES[data['period']]} series", styles["Heading3"]),
        st,
        Spacer(1, 8),
    ]

    plants = [
        ["Plant", "Saved (INR)", "Energy saved (kWh)", "Feedstock saved (kg)", "Batches rescued"]
    ] + [
        [
            p["name"],
            _fmt(p["money_saved_inr"]),
            _fmt(p["energy_saved_kwh"]),
            _fmt(p["feedstock_saved_kg"]),
            str(p["batches_rescued"]),
        ]
        for p in data["by_plant"]
    ]
    pt = Table(plants)
    pt.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), teal),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
            ]
        )
    )
    story += [Paragraph("By plant (12 months with Kairos)", styles["Heading3"]), pt, Spacer(1, 8)]

    asm = [["Assumption", "Value", "Unit", "Provenance"]] + [
        [Paragraph(a["label"], small), str(a["value"]), a["unit"], a["provenance"]]
        for a in data["provenance"]["assumptions"]
    ]
    at = Table(asm, colWidths=[120 * mm, 30 * mm, 25 * mm, 25 * mm])
    at.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), teal),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story += [
        Paragraph("Assumptions", styles["Heading3"]),
        at,
        Spacer(1, 6),
        Paragraph(data["provenance"]["simulated"], note),
    ]
    doc.build(story)
    return buf.getvalue()
